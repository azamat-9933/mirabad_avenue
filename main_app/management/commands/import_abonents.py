"""Import houses, apartments, personal accounts, areas and owners from the abonent registry.

The registry is a folder with one workbook per house, named after the house (60.xlsx, 60V.xlsx,
...). Each sheet has the columns "Abonent turi", "Abonent raqami", "F.I.O.", "Manzil",
"Kadastr raqami" and "Umumiy maydon". Apartments are matched by their personal account first and
by house and number second, so running the command again updates them in place. Balances,
phones, Telegram links and billing data are never touched.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from main_app.models import DEFAULT_SECTOR_NAME, Apartment, Building, Complex, Owner


AREA_PLACES = Decimal("0.01")
DEFAULT_AUTHOR = "Abonent registry"
NON_RESIDENTIAL_PREFIX = "неж"

HEADERS = {
    "kind": "abonent turi",
    "account": "abonent raqami",
    "fio": "f.i.o.",
    "address": "manzil",
    "area": "umumiy maydon",
}
NON_RESIDENTIAL_KINDS = {"noturar joy", "noturarjoy"}

# Letters after an apartment number ("24а", "83A") are written in Latin capitals.
LETTERS = str.maketrans({"а": "A", "А": "A", "a": "A", "б": "B", "Б": "B", "b": "B", "в": "V", "В": "V", "v": "V"})
NUMBER_PATTERNS = (
    re.compile(r"кв\.?\s*(\d+\s*[^\W\d_]?)\s*$", re.IGNORECASE),
    re.compile(r"(\d+\s*[^\W\d_]?)\s*-\s*xonadon\s*$", re.IGNORECASE),
    # 64V writes the apartment right after the house: "ул. Миробод, 64-V, 12".
    re.compile(r"\d+\s*-\s*V\s*,\s*(\d+\s*[^\W\d_]?)\s*$", re.IGNORECASE),
)
NON_RESIDENTIAL_PATTERN = re.compile(r"\bнеж[а-яё]*\.?\s*(\d*)\s*$", re.IGNORECASE)
HOUSE_IN_ADDRESS = re.compile(r"(?:дом\s*|\b)(\d+)\s*-?\s*([vв])?\s*(?:-uy|,|$)", re.IGNORECASE)


@dataclass
class RegistryRow:
    house: str
    row: int
    account: str
    fio: str
    address: str
    area: Decimal | None
    non_residential: bool
    number: str = ""
    notes: list[str] = field(default_factory=list)


def clean_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\t", " ").split()).strip()


def clean_account(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value).replace(" ", "")


def to_area(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", ".")).quantize(AREA_PLACES, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def apartment_number(address: str) -> str:
    """'ул. Миробод, дом 62, кв. 24a' -> '24A'; '..., 66 V-uy, 137A-xonadon' -> '137A'."""
    for pattern in NUMBER_PATTERNS:
        match = pattern.search(address)
        if match:
            return match.group(1).replace(" ", "").translate(LETTERS)
    return ""


def non_residential_number(address: str) -> str | None:
    """'..., неж3' -> '3'; '..., неж' or '..., нежелой' -> ''; no such word -> None."""
    match = NON_RESIDENTIAL_PATTERN.search(address)
    return match.group(1) if match else None


def house_in_address(address: str) -> str:
    """'ул. Миробод, дом 60в, неж' -> '60V'; "Mirobod ko'chasi, 64v-uy" -> '64V'."""
    match = HOUSE_IN_ADDRESS.search(address)
    return match.group(1) + ("V" if match.group(2) else "") if match else ""


def read_house(path: Path) -> list[RegistryRow]:
    house = path.stem.strip().upper()
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    header = [clean_text(cell).lower() for cell in rows[0]]
    columns = {key: header.index(title) for key, title in HEADERS.items() if title in header}
    missing = [HEADERS[key] for key in ("account", "fio", "address", "area") if key not in columns]
    if missing:
        raise CommandError(f"{path.name}: no column {', '.join(missing)}")

    def cell(values, key):
        index = columns.get(key)
        return values[index] if index is not None and index < len(values) else None

    result = []
    for row_number, values in enumerate(rows[1:], start=2):
        account = clean_account(cell(values, "account"))
        if not account:
            continue  # the total line under the header, empty lines
        address = clean_text(cell(values, "address"))
        kind = clean_text(cell(values, "kind")).lower()
        result.append(
            RegistryRow(
                house=house,
                row=row_number,
                account=account,
                fio=clean_text(cell(values, "fio")),
                address=address,
                area=to_area(cell(values, "area")),
                non_residential=kind in NON_RESIDENTIAL_KINDS or non_residential_number(address) is not None,
            )
        )
    return result


def number_units(rows: list[RegistryRow]) -> list[str]:
    """Give every row of one house its apartment number; return the problems found."""
    problems = []
    unnamed = []
    for row in rows:
        if row.non_residential:
            digits = non_residential_number(row.address)
            if digits:
                row.number = f"{NON_RESIDENTIAL_PREFIX}{digits}"
            else:
                unnamed.append(row)
            continue
        row.number = apartment_number(row.address)
        if not row.number:
            problems.append(f"{row.house}.xlsx row {row.row}: no apartment number in the address {row.address!r}")

    # Non-residential units the address does not number: "неж" if it is the only one, otherwise
    # неж1, неж2, ... in the order of the file, skipping numbers already used in the house.
    if len(unnamed) == 1:
        unnamed[0].number = NON_RESIDENTIAL_PREFIX
    else:
        used = {row.number for row in rows if row.number}
        counter = 0
        for row in unnamed:
            counter += 1
            while f"{NON_RESIDENTIAL_PREFIX}{counter}" in used:
                counter += 1
            row.number = f"{NON_RESIDENTIAL_PREFIX}{counter}"
            used.add(row.number)
    for row in unnamed:
        row.notes.append(f"the address gives no unit number, saved as {row.number}")
    return problems


def check_registry(houses: dict[str, list[RegistryRow]]) -> list[str]:
    """Number the units and return every problem that stops the import."""
    problems = []
    seen_accounts: dict[str, RegistryRow] = {}
    for house, rows in houses.items():
        problems.extend(number_units(rows))
        seen_numbers: dict[str, RegistryRow] = {}
        for row in rows:
            if not row.account.isdigit():
                problems.append(f"{house}.xlsx row {row.row}: the personal account {row.account!r} is not a number")
            if row.area is None or row.area < 0:
                problems.append(f"{house}.xlsx row {row.row}: no area")
            elif row.area == 0:
                row.notes.append("the area is 0")
            if not row.fio:
                row.notes.append("no owner name: the apartment stays vacant")
            named = house_in_address(row.address)
            if named and named != house:
                row.notes.append(f"the address names house {named}: {row.address}")
            first = seen_accounts.setdefault(row.account, row)
            if first is not row:
                problems.append(
                    f"personal account {row.account} is on {first.house}.xlsx row {first.row} and {house}.xlsx row {row.row}"
                )
            if row.number:
                twin = seen_numbers.setdefault(row.number, row)
                if twin is not row:
                    problems.append(f"{house}.xlsx rows {twin.row} and {row.row}: both are apartment {row.number}")
    return problems


def building_address(rows: list[RegistryRow]) -> str:
    """The most common address of the house's apartments, without the apartment part."""
    counts: dict[str, int] = {}
    for row in rows:
        if row.non_residential:
            continue
        base = re.sub(r",?\s*(?:кв\.?\s*)?\d+\s*[^\W\d_]?\s*(?:-\s*xonadon)?\s*$", "", row.address, flags=re.IGNORECASE)
        base = base.strip(" ,")
        if base:
            counts[base] = counts.get(base, 0) + 1
    return max(counts, key=counts.get) if counts else ""


class Command(BaseCommand):
    help = "Import houses, apartments, personal accounts, areas and owners from the abonent registry folder."

    def add_arguments(self, parser):
        parser.add_argument("folder", help="Folder with one workbook per house: 60.xlsx, 60V.xlsx, ...")
        parser.add_argument("--author", default=DEFAULT_AUTHOR, help=f"Author of created rows. Default: {DEFAULT_AUTHOR}")
        parser.add_argument("--dry-run", action="store_true", help="Check the files and report without writing anything.")

    def handle(self, *args, **options):
        folder = Path(options["folder"])
        files = sorted(path for path in folder.glob("*.xlsx") if not path.name.startswith("~$"))
        if not files:
            raise CommandError(f"No .xlsx files in {folder}")

        houses = {path.stem.strip().upper(): read_house(path) for path in files}
        problems = check_registry(houses)
        if problems:
            raise CommandError("The registry has problems, nothing was imported:\n" + "\n".join(problems))

        for house, rows in houses.items():
            living = sum(1 for row in rows if not row.non_residential)
            area = sum((row.area for row in rows), Decimal("0"))
            self.stdout.write(
                f"House {house}: {len(rows)} units ({living} apartments, {len(rows) - living} non-residential), "
                f"area {area} m²"
            )
            for row in rows:
                for note in row.notes:
                    self.stdout.write(f"  {house}.xlsx row {row.row}, {row.number} ({row.account}): {note}")

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run: the database was not changed."))
            return

        stats = dict.fromkeys(
            ("buildings_created", "apartments_created", "apartments_updated", "owners_created", "owners_updated"), 0
        )
        with transaction.atomic():
            sector = Complex.objects.order_by("pk").first() or Complex.objects.create(
                title=DEFAULT_SECTOR_NAME, address=DEFAULT_SECTOR_NAME, author=options["author"]
            )
            for house, rows in houses.items():
                building = Building.objects.filter(number__iexact=house).order_by("pk").first()
                if building is None:
                    building = Building.objects.create(
                        complex=sector, number=house, address=building_address(rows), author=options["author"]
                    )
                    stats["buildings_created"] += 1
                elif not building.address:
                    building.address = building_address(rows)
                    building.save(update_fields=["address"])
                for row in rows:
                    self._import_row(building, row, options["author"], stats)

        for key, value in stats.items():
            self.stdout.write(f"{key}: {value}")
        self.stdout.write(self.style.SUCCESS("The abonent registry was imported."))

    def _import_row(self, building, row: RegistryRow, author: str, stats: dict):
        apartment = (
            Apartment.objects.filter(personal_account=row.account).first()
            or Apartment.objects.filter(building=building, number=row.number).order_by("pk").first()
        )
        if apartment is None:
            apartment = Apartment.objects.create(
                building=building, number=row.number, area=row.area, personal_account=row.account, author=author
            )
            stats["apartments_created"] += 1
        else:
            values = {"building": building, "number": row.number, "area": row.area, "personal_account": row.account}
            changed = [name for name, value in values.items() if getattr(apartment, name) != value]
            for name in changed:
                setattr(apartment, name, values[name])
            if changed:
                # Only these columns: the balance changes through the ledger alone.
                apartment.save(update_fields=changed)
                stats["apartments_updated"] += 1

        if not row.fio:
            return
        owner = Owner.objects.filter(apartment=apartment).first()
        if owner is None:
            # A new owner takes over the apartment account as it is (Owner.save copies the balance).
            Owner.objects.create(apartment=apartment, fio=row.fio, phone="")
            stats["owners_created"] += 1
        elif owner.fio != row.fio:
            owner.fio = row.fio
            owner.save(update_fields=["fio"])
            stats["owners_updated"] += 1
