from datetime import date
import base64
import gzip
import io
import json
import re
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.db import connection
from django.test import Client, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from billing.models import BillingPeriod, Charge, Expense, HeatingRecord, HotWaterMeterReading
from main_app.models import Apartment, Building, Complex, Owner
from payments.models import Transaction
from payments.services import OPENING_BALANCE_TAG, apply_account_change, apply_balance_change
from portal.audit import record_event
from portal.backend_data import build_portal_data
from portal.models import AuditEvent


PORTAL_PAGES = [
    "portal:dashboard",
    "portal:residential",
    "portal:billing",
    "portal:billing_periods",
    "portal:analytics",
]


class PortalAuthenticationTests(TestCase):
    password = "portal-test-password"

    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.user = User.objects.create_user(username="operator", password=cls.password)
        cls.superuser = User.objects.create_superuser(username="root-admin", password=cls.password)
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64", complex=complex_obj, author="test")
        apartment = Apartment.objects.create(number="1", area=Decimal("50.00"), building=building, author="test")
        cls.owner = Owner.objects.create(fio="Test Owner", phone="+998901234567", apartment=apartment)

    def _payment_payload(self):
        return {"owner_id": self.owner.id, "amount": "1000", "operation": "credit", "channel": "cash"}

    def test_portal_pages_redirect_anonymous_users_to_login(self):
        login_url = reverse("portal:login")
        for name in PORTAL_PAGES:
            with self.subTest(page=name):
                path = reverse(name)
                response = self.client.get(path)
                self.assertRedirects(response, f"{login_url}?next={path}", fetch_redirect_response=False)

    def test_api_returns_401_for_anonymous_users(self):
        for name in [
            "portal:api_root",
            "portal:api_portal_data",
            "portal:api_portal_data_version",
            "portal:api_list_residents",
            "portal:api_list_transactions",
            "portal:api_list_audit",
        ]:
            with self.subTest(endpoint=name):
                response = self.client.get(reverse(name), HTTP_ACCEPT="application/json")
                self.assertEqual(response.status_code, 401)

    def test_anonymous_user_cannot_create_payment(self):
        response = self.client.post(
            reverse("portal:api_create_payment"),
            self._payment_payload(),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 401)
        self.owner.refresh_from_db()
        self.assertEqual(self.owner.balance, Decimal("0"))
        self.assertFalse(Transaction.objects.exists())

    def test_health_endpoint_stays_public(self):
        response = self.client.get(reverse("portal:api_health"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok"})

    def test_api_schema_and_docs_require_login(self):
        for path in ["/api/schema/", "/api/docs/"]:
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertRedirects(
                    response,
                    f"{reverse('portal:login')}?next={path}",
                    fetch_redirect_response=False,
                )

    def test_login_page_renders_for_anonymous_users(self):
        response = self.client.get(reverse("portal:login"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "portal/login.html")

    def test_login_redirects_to_next(self):
        response = self.client.post(
            reverse("portal:login"),
            {"username": "operator", "password": self.password, "next": "/billing/"},
        )
        self.assertRedirects(response, "/billing/", fetch_redirect_response=False)

    def test_login_ignores_external_next(self):
        response = self.client.post(
            reverse("portal:login"),
            {"username": "operator", "password": self.password, "next": "https://evil.example/"},
        )
        self.assertRedirects(response, reverse("portal:dashboard"), fetch_redirect_response=False)

    def test_wrong_password_shows_error(self):
        response = self.client.post(reverse("portal:login"), {"username": "operator", "password": "wrong"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].errors)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_inactive_user_cannot_log_in(self):
        self.user.is_active = False
        self.user.save(update_fields=["is_active"])
        response = self.client.post(reverse("portal:login"), {"username": "operator", "password": self.password})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_authenticated_user_can_open_pages_and_api(self):
        self.client.force_login(self.user)
        for name in PORTAL_PAGES:
            with self.subTest(page=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        response = self.client.get(reverse("portal:api_portal_data"), HTTP_ACCEPT="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["profile"]["username"], "operator")

    def test_authenticated_post_requires_csrf_token(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        url = reverse("portal:api_create_payment")

        response = client.post(url, self._payment_payload(), content_type="application/json")
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Transaction.objects.exists())

        client.get(reverse("portal:dashboard"))
        response = client.post(
            url,
            self._payment_payload(),
            content_type="application/json",
            HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value,
        )
        self.assertEqual(response.status_code, 200)
        self.owner.refresh_from_db()
        self.assertEqual(self.owner.balance, Decimal("1000"))

    def test_operators_take_cash_card_or_card_transfer_only(self):
        # Payme and Click payments come from those services' integrations, not from the portal form.
        self.client.force_login(self.user)
        url = reverse("portal:api_create_payment")
        for channel in ("payme", "click"):
            with self.subTest(channel=channel):
                response = self.client.post(
                    url, {**self._payment_payload(), "channel": channel}, content_type="application/json"
                )
                self.assertEqual(response.status_code, 400)
        self.assertFalse(Transaction.objects.exists())

        for channel, tag in (("card", "[portal:card-credit]"), ("transfer", "[portal:transfer-credit]")):
            with self.subTest(channel=channel):
                response = self.client.post(
                    url, {**self._payment_payload(), "channel": channel}, content_type="application/json"
                )
                self.assertEqual(response.status_code, 200)
                self.assertIn(tag, Transaction.objects.latest("id").description)

    def test_card_transfer_is_its_own_payment_channel(self):
        self.client.force_login(self.user)
        self.client.post(
            reverse("portal:api_create_payment"),
            {**self._payment_payload(), "amount": "250000", "channel": "transfer"},
            content_type="application/json",
        )

        snapshot = build_portal_data(self.user)
        [row] = snapshot["transactions"]
        self.assertEqual((row["method"], row["type"]), ("Card transfer", "Utility Payment"))
        self.assertEqual(snapshot["paymentSummary"]["channels"]["Card transfer"]["amount"], 250000.0)
        # Counted as money received: it becomes the resident's last payment.
        resident = next(row for row in snapshot["residents"] if row["backendId"] == self.owner.id)
        self.assertTrue(resident["lastPayment"])

    def test_logout_ends_session_and_points_to_login(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("portal:api_logout"), {}, content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["redirectUrl"], reverse("portal:login"))
        self.assertNotIn("_auth_user_id", self.client.session)
        response = self.client.get(reverse("portal:dashboard"))
        self.assertEqual(response.status_code, 302)

    def test_logout_no_longer_accepts_get_redirects(self):
        response = self.client.get(reverse("portal:api_logout"), {"next": "https://evil.example/"})
        self.assertEqual(response.status_code, 405)

    def test_profile_does_not_fall_back_to_another_account(self):
        profile = build_portal_data(None)["profile"]
        self.assertEqual(profile["username"], "")
        self.assertFalse(profile["isAuthenticated"])

    def test_admin_login_stays_reachable_and_admin_requires_staff(self):
        self.assertEqual(self.client.get(reverse("admin:login")).status_code, 200)
        # Every operator account has full rights (portal.accounts); one narrowed by hand loses admin.
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("admin:index")).status_code, 200)
        type(self.user).objects.filter(pk=self.user.pk).update(is_staff=False, is_superuser=False)
        response = self.client.get(reverse("admin:index"))
        self.assertRedirects(
            response,
            f"{reverse('admin:login')}?next={reverse('admin:index')}",
            fetch_redirect_response=False,
        )


class PortalAssetTests(TestCase):
    """The page loads built CSS and only the icon glyphs it uses, not the Tailwind Play CDN."""

    def test_base_template_uses_the_built_tailwind_css(self):
        base = (Path(settings.BASE_DIR) / "templates" / "base.html").read_text(encoding="utf-8")
        self.assertNotIn("cdn.tailwindcss.com", base)
        self.assertIn("css/tailwind.css", base)
        self.assertTrue((Path(settings.BASE_DIR) / "static" / "css" / "tailwind.css").is_file())

    def test_every_icon_written_in_templates_and_scripts_is_in_the_icon_font(self):
        root = Path(settings.BASE_DIR)
        base = (root / "templates" / "base.html").read_text(encoding="utf-8")
        names = re.search(r"icon_names=([a-z0-9_,]+)", base).group(1).split(",")
        self.assertEqual(names, sorted(names))  # Google Fonts wants them in order
        used = set()
        sources = list((root / "templates").glob("**/*.html")) + list((root / "static" / "js").glob("*.js"))
        for path in sources:
            text = path.read_text(encoding="utf-8")
            used.update(re.findall(r'material-symbols-outlined[^"]*"[^>]*>\s*([a-z][a-z0-9_]*)\s*<', text))
        self.assertTrue(used)
        self.assertEqual(sorted(used - set(names)), [])


class PortalSnapshotTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64", complex=complex_obj, author="test")
        apartment = Apartment.objects.create(number="1", area=Decimal("50.00"), building=building, author="test")
        cls.owner = Owner.objects.create(
            fio="Иванов </script><script>alert(1)</script> & <b>",
            phone="+998901234567",
            apartment=apartment,
        )

    def setUp(self):
        self.client.force_login(self.user)

    def _embedded_data(self, html):
        match = re.search(r'<script id="hydroflow-backend-data" type="application/json">(.*?)</script>', html, re.S)
        self.assertIsNotNone(match)
        return match.group(1)

    def test_page_embeds_snapshot_that_cannot_close_the_script_tag(self):
        html = self.client.get(reverse("portal:dashboard")).content.decode("utf-8")
        raw = self._embedded_data(html)

        self.assertNotIn("<", raw)
        self.assertIn("Иванов", raw)  # non-ASCII stays readable instead of \u-escaped
        data = json.loads(raw)
        self.assertEqual(data["residents"][0]["name"], self.owner.fio)

    def test_pages_poll_a_version_that_changes_with_the_data(self):
        url = reverse("portal:api_portal_data_version")
        version = self.client.get(url).json()["version"]
        self.assertEqual(self.client.get(reverse("portal:api_portal_data")).json()["dataVersion"], version)
        self.assertEqual(self.client.get(url).json()["version"], version)  # nothing changed

        apply_balance_change(self.owner.id, Decimal("1000"), Transaction.TYPE_CASH)
        after_payment = self.client.get(url).json()["version"]
        self.assertNotEqual(after_payment, version)
        # Edits that leave no ledger or log row are seen too.
        Owner.objects.filter(pk=self.owner.pk).update(phone="+998900000000")
        self.assertNotEqual(self.client.get(url).json()["version"], after_payment)

    def test_snapshot_is_reused_until_the_data_changes(self):
        first = build_portal_data(self.user)
        with CaptureQueriesContext(connection) as queries:
            again = build_portal_data(self.user)
        # Only the version, the profile and the audit rows are read again.
        self.assertLess(len(queries), 30)
        self.assertEqual(again["residents"], first["residents"])

        apply_balance_change(self.owner.id, Decimal("1000"), Transaction.TYPE_CASH)
        self.assertEqual(build_portal_data(self.user)["residents"][0]["balance"], 1000.0)

    def test_dashboard_lists_residents_a_page_at_a_time_with_a_search(self):
        page = self.client.get(reverse("portal:dashboard")).content.decode("utf-8")
        self.assertIn('data-resident-paged="true"', page)
        self.assertIn('data-resident-page-size="50"', page)
        self.assertIn("data-resident-pager", page)
        self.assertIn("data-resident-search", page)
        self.assertNotIn('data-resident-fetch-all="true"', page)

    def test_responses_are_gzipped_when_the_client_accepts_it(self):
        response = self.client.get(reverse("portal:api_portal_data"), HTTP_ACCEPT_ENCODING="gzip")
        self.assertEqual(response["Content-Encoding"], "gzip")
        self.assertEqual(json.loads(gzip.decompress(response.content))["residents"][0]["backendId"], self.owner.id)

    def test_complexes_list_returns_rows(self):
        apply_balance_change(self.owner.id, Decimal("1000"), Transaction.TYPE_CASH)
        response = self.client.get(reverse("portal:api_list_complexes"))
        self.assertEqual(response.status_code, 200)
        row = response.json()["results"][0]
        self.assertEqual(row["units"], 1)
        self.assertEqual(row["baselineCollected"], 1000.0)

    def test_transaction_rows_keep_their_labels(self):
        flavours = [
            (Transaction.TYPE_CASH, "1000", "", "Utility Payment", "Cash", "Success"),
            (Transaction.TYPE_MANUAL, "500", "[portal:card-credit] Portal credit via card.", "Utility Payment", "Card", "Success"),
            (Transaction.TYPE_MANUAL, "-200", "[portal:balance-debit] x", "Manual Adjustment", "Adjustment", "Pending"),
            (Transaction.TYPE_MANUAL, "-300", f"{OPENING_BALANCE_TAG} x", "Manual Adjustment", "Manual", "Pending"),
            (Transaction.TYPE_CHARGE, "-400", "[billing:period-charge] x", "Monthly Charge", "Charge", "Pending"),
            (Transaction.TYPE_CHARGE, "100", "[billing:recalculation] x", "Recalculation", "Charge", "Pending"),
        ]
        for payment_type, amount, description, *_ in flavours:
            apply_balance_change(self.owner.id, Decimal(amount), payment_type, description=description)

        listed = self.client.get(reverse("portal:api_list_transactions"), {"page_size": 100}).json()["results"]
        rows = {row["backendId"]: row for row in listed}
        ledger = Transaction.objects.filter(owner=self.owner).order_by("id")
        for item, (_, amount, _, expected_type, expected_method, expected_status) in zip(ledger, flavours):
            with self.subTest(description=item.description):
                row = rows[item.id]
                self.assertEqual(
                    (row["type"], row["method"], row["status"], row["signedAmount"]),
                    (expected_type, expected_method, expected_status, float(amount)),
                )
                self.assertRegex(row["date"], r"^\d{2}\.\d{2}\.\d{4}$")
                self.assertRegex(row["createdAt"], r"^\d{2}\.\d{2}\.\d{4}, \d{2}:\d{2}$")


    def test_snapshot_and_list_endpoints_describe_rows_identically(self):
        Owner.objects.filter(pk=self.owner.pk).update(telegram_status=Owner.TELEGRAM_CONNECTED, has_contract=True)
        apply_balance_change(self.owner.id, Decimal("1000"), Transaction.TYPE_CASH, description="[portal:cash-credit] x")
        snapshot = build_portal_data(self.user)

        listed_resident = self.client.get(reverse("portal:api_list_residents"), {"owner_id": self.owner.id}).json()["results"][0]
        self.assertEqual(listed_resident, snapshot["residents"][0])
        self.assertEqual(listed_resident["telegramStatus"], "connected")

        listed_transaction = self.client.get(reverse("portal:api_list_transactions")).json()["results"][0]
        self.assertEqual(listed_transaction, snapshot["transactions"][0])

class PortalExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64", complex=complex_obj, author="test")
        apartment = Apartment.objects.create(number="1", area=Decimal("50.00"), building=building, author="test")
        cls.owner = Owner.objects.create(fio="Ўринбоев Ғайрат", phone="+998901234567", apartment=apartment)
        for index in range(120):
            apply_balance_change(cls.owner.id, Decimal(1000 + index), Transaction.TYPE_CASH)

    def setUp(self):
        self.client.force_login(self.user)

    def _export(self, url, payload):
        response = self.client.post(url, payload, content_type="application/json")
        self.assertEqual(response.status_code, 200)
        download = response.json()["download"]
        return download, base64.b64decode(download["content"])

    def test_transactions_xlsx_has_all_rows_and_localized_headers(self):
        download, content = self._export(
            reverse("portal:api_export_report"),
            {"context": "transactions", "format": "XLSX", "language": "ru"},
        )
        self.assertTrue(download["filename"].endswith(".xlsx"))
        sheet = load_workbook(io.BytesIO(content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            ("ID", "Дата и время", "Житель", "Оператор", "Тип оплаты", "Сумма", "Баланс до", "Баланс после", "Статус"),
        )
        self.assertEqual(len(rows), 121)
        self.assertEqual(rows[1][2], "Ўринбоев Ғайрат")

    def test_pdf_export_is_not_truncated(self):
        download, content = self._export(
            reverse("portal:api_export_report"),
            {"context": "transactions", "format": "PDF", "language": "uz"},
        )
        self.assertEqual(download["mime"], "application/pdf")
        self.assertTrue(content.startswith(b"%PDF"))
        # 120 rows cannot fit on one page; the old writer silently stopped after 45 lines.
        self.assertGreater(len(re.findall(rb"/Type /Page\b", content)), 1)

    def test_residential_exports_render(self):
        for export_format in ["XLSX", "PDF"]:
            with self.subTest(export_format=export_format):
                download, content = self._export(
                    reverse("portal:api_export_report"),
                    {"context": "residential", "format": export_format},
                )
                self.assertTrue(download["filename"].startswith("residential"))
                self.assertTrue(content)

    def test_removed_maintenance_context_falls_back_to_residential(self):
        download, _content = self._export(
            reverse("portal:api_export_report"),
            {"context": "maintenance", "format": "XLSX"},
        )
        self.assertTrue(download["filename"].startswith("residential"))

    def test_audit_export_is_xlsx(self):
        self.client.force_login(get_user_model().objects.create_superuser(username="head-admin", password="unused"))
        download, content = self._export(reverse("portal:api_export_audit"), {"language": "uz"})
        self.assertTrue(download["filename"].startswith("audit-log-"))
        header = next(load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True))
        self.assertEqual(header[:3], ("Audit ID", "Sana va vaqt", "Hodisa turi"))


class PortalDebtReportTests(TestCase):
    """Export window reports: debtors list, payment history and both together, grouped by house."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        cls.house_64 = Building.objects.create(number="64A", complex=complex_obj, author="test")
        cls.house_12 = Building.objects.create(number="12", complex=complex_obj, author="test")
        debtor_flat = Apartment.objects.create(number="2", area=Decimal("76.50"), building=cls.house_64, author="test")
        paid_flat = Apartment.objects.create(number="1", area=Decimal("64.50"), building=cls.house_64, author="test")
        cls.vacant_flat = Apartment.objects.create(number="9", area=Decimal("40"), building=cls.house_12, author="test")
        cls.debtor = Owner.objects.create(fio="Muhammad Yaxyoyev", phone="+998773487001", apartment=debtor_flat)
        cls.payer = Owner.objects.create(fio="Azamat Yaxyoyev", phone="+998901139933", apartment=paid_flat)
        apply_balance_change(cls.debtor.id, Decimal("-300000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        apply_balance_change(cls.payer.id, Decimal("-100000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        apply_balance_change(
            cls.payer.id, Decimal("400000"), Transaction.TYPE_MANUAL,
            description="[portal:transfer-credit] Portal credit via card transfer.", created_by="desk",
        )
        apply_account_change(cls.vacant_flat.id, Decimal("-80000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")

    def setUp(self):
        self.client.force_login(self.user)

    def _workbook(self, payload):
        response = self.client.post(
            reverse("portal:api_export_report"), {"format": "XLSX", "language": "ru", **payload}, content_type="application/json"
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(base64.b64decode(response.json()["download"]["content"])))
        return {
            sheet.title: [row for row in sheet.iter_rows(values_only=True) if any(value not in (None, "") for value in row)]
            for sheet in workbook
        }

    def test_debtors_list_groups_debtors_by_house(self):
        rows = self._workbook({"report_type": "debtors"})["Список должников"]
        self.assertEqual(rows[3], ("ID", "Дата", "Ф.И.О.", "Дом", "Квартира", "Площадь (м2)", "Сумма долга (UZS)"))
        today = timezone.localdate().strftime("%d.%m.%Y")
        self.assertEqual(rows[4][0], "Дом 12")
        # The vacant apartment's debt is the apartment's; the resident who paid is not listed.
        self.assertEqual(rows[5], ("—", today, "Владелец не назначен", "12", "9", 40, 80000))
        self.assertEqual(rows[6][0], "Итого по дому 12 · Должников: 1")
        self.assertEqual(rows[7][0], "Дом 64A")
        self.assertEqual(rows[8], (str(self.debtor.pk), today, "Muhammad Yaxyoyev", "64A", "2", 76.5, 300000))
        self.assertEqual(rows[10][0], "Итого · Должников: 2")
        self.assertEqual(rows[10][5:], (116.5, 380000))

    def test_selected_house_limits_the_report(self):
        rows = self._workbook({"report_type": "debtors", "building_backend_ids": [self.house_64.id]})["Список должников"]
        self.assertEqual(rows[1][0], "Дома: 64A")
        self.assertEqual([row[2] for row in rows[5:-2]], ["Muhammad Yaxyoyev"])

    def test_debt_is_taken_as_of_date_to(self):
        # The resident owed 500 last week and paid today: listed for last week, not for today.
        flat = Apartment.objects.create(number="3", area=Decimal("50"), building=self.house_64, author="test")
        late = Owner.objects.create(fio="Late Payer", phone="+998901000003", apartment=flat)
        charge = apply_balance_change(late.id, Decimal("-500"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        Transaction.objects.filter(pk=charge.pk).update(created_at=timezone.now() - timedelta(days=7))
        apply_balance_change(late.id, Decimal("500"), Transaction.TYPE_CASH)

        names_now = [row[2] for row in self._workbook({"report_type": "debtors"})["Список должников"]]
        self.assertNotIn("Late Payer", names_now)
        last_week = (timezone.localtime() - timedelta(days=3)).strftime("%Y-%m-%dT%H:%M")
        rows = self._workbook({"report_type": "debtors", "date_to": last_week})["Список должников"]
        self.assertIn(("Late Payer", 500), [(row[2], row[6]) for row in rows])

    def test_payment_history_lists_payments_and_subtractions(self):
        apply_balance_change(
            self.payer.id, Decimal("-150000"), Transaction.TYPE_MANUAL,
            description="[portal:balance-debit] Portal manual balance subtraction.", created_by="desk",
        )
        rows = self._workbook({"report_type": "payment_history"})["История платежей"]
        header = rows.index(("ID", "Дата и время", "Ф.И.О.", "Квартира", "Способ оплаты", "Сумма (UZS)", "Оператор"))
        body = rows[header + 1:]
        # Billing charges are left out; the transfer and the subtraction are listed under their house.
        self.assertEqual(body[0][0], "Дом 64A")
        self.assertEqual(body[1][2:], ("Azamat Yaxyoyev", "1", "Перевод на карту", 400000, "desk"))
        self.assertEqual(body[2][2:], ("Azamat Yaxyoyev", "1", "Вычитание баланса", -150000, "desk"))
        self.assertEqual(body[-1][0], "Итого · Платежей: 1 (400 000.00) · Вычетов: 1 (-150 000.00)")
        self.assertEqual(body[-1][5], 250000)

    def test_payment_history_respects_the_date_range(self):
        tomorrow = (timezone.localtime() + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")
        rows = self._workbook({"report_type": "payment_history", "date_from": tomorrow})["История платежей"]
        self.assertEqual(rows[-1], ("Нет строк для текущих фильтров.", None, None, None, None, None, None))

    def test_general_overview_is_one_workbook_with_both_lists(self):
        sheets = self._workbook({"report_type": "general"})
        self.assertEqual(list(sheets), ["Список должников", "История платежей"])

    def test_reports_render_as_pdf(self):
        for report_type in ["general", "debtors", "payment_history"]:
            with self.subTest(report_type=report_type):
                response = self.client.post(
                    reverse("portal:api_export_report"),
                    {"report_type": report_type, "format": "PDF", "language": "uz"},
                    content_type="application/json",
                )
                download = response.json()["download"]
                self.assertEqual(download["mime"], "application/pdf")
                self.assertTrue(base64.b64decode(download["content"]).startswith(b"%PDF"))

    def test_apartment_balance_history_has_every_ledger_entry(self):
        sheets = self._workbook({"report_type": "account_history", "apartment_backend_ids": [self.payer.apartment_id]})
        rows = sheets["История баланса"]
        self.assertEqual(rows[0][0], "История баланса: Дом 64A, Квартира 1")
        self.assertEqual(rows[1][0], "Владелец: Azamat Yaxyoyev")
        header = rows.index(("ID", "Дата и время", "Абонент", "Операция", "Оператор", "Сумма (UZS)", "Баланс до", "Баланс после"))
        # Oldest first: the charge, then the transfer; the balance runs from one entry to the next.
        self.assertEqual(
            [row[3:] for row in rows[header + 1:]],
            [
                ("Начисление за месяц", "Система", -100000, 0, -100000),
                ("Перевод на карту", "desk", 400000, -100000, 300000),
            ],
        )

    def test_balance_history_needs_an_apartment(self):
        response = self.client.post(
            reverse("portal:api_export_report"), {"report_type": "account_history"}, content_type="application/json"
        )
        self.assertEqual(response.status_code, 400)

    def test_apartment_window_gets_the_owners_telegram(self):
        Owner.objects.filter(pk=self.payer.pk).update(
            telegram_status=Owner.TELEGRAM_CONNECTED, telegram_id="77", telegram_user="azamat_y"
        )
        building = next(
            item for item in build_portal_data(self.user)["complexes"][0]["buildingItems"] if item["number"] == "64A"
        )
        owner = next(item["owner"] for item in building["apartments"] if item["unit"] == "1")
        self.assertEqual((owner["telegramUser"], owner["telegramConnected"]), ("azamat_y", True))

    def test_houses_page_has_no_recent_alerts_panel(self):
        content = self.client.get(reverse("portal:residential")).content.decode()
        self.assertNotIn("recent-alerts", content)
        self.assertIn('data-report-type="debtors"', content)


DEMO_MARKERS = [
    "Harbor View",
    "Sarah Jenkins",
    "Marcus Aurelius",
    "Aziz Karimov",
    "HF-9284",
    "images.unsplash.com",
    "Balance report exported by Admin User",
    "@hydroflow.local",
]


class PortalDemoContentTests(TestCase):
    """Pages and the frontend bundle must show only backend data, never invented sample records."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")

    def test_pages_render_without_sample_records(self):
        self.client.force_login(self.user)
        for name in PORTAL_PAGES:
            html = self.client.get(reverse(name)).content.decode("utf-8")
            for marker in [*DEMO_MARKERS, "Skyline"]:
                with self.subTest(page=name, marker=marker):
                    self.assertFalse(marker in html, f"{marker!r} rendered on {name}")

    def test_frontend_bundle_has_no_sample_dataset(self):
        bundle = (Path(settings.BASE_DIR) / "static" / "js" / "app.js").read_text(encoding="utf-8")
        for marker in [*DEMO_MARKERS, "buildAnalyticsDemoData", "HydroFlowMockData", "Marcus"]:
            with self.subTest(marker=marker):
                self.assertFalse(marker in bundle, f"{marker!r} found in app.js")


class PortalRemovedSystemHealthTests(TestCase):
    """The System Health page and its maintenance/telemetry backend were removed and must stay gone."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")

    def test_system_health_urls_are_gone(self):
        self.client.force_login(self.user)
        for path in ["/system-health/", "/api/lists/maintenance/", "/api/maintenance/deploy/"]:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 404)

    def test_snapshot_has_no_maintenance_or_telemetry(self):
        data = build_portal_data(self.user)
        for key in ["maintenanceTasks", "telemetryNodes", "pressureSeries"]:
            with self.subTest(key=key):
                self.assertNotIn(key, data)


class PortalApartmentAccountTests(TestCase):
    """Debt billed to a vacant apartment shows up in the portal and passes to its next owner."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="66", complex=complex_obj, author="test")
        cls.apartment = Apartment.objects.create(number="7", area=Decimal("60.00"), building=building, author="test")
        apply_account_change(
            cls.apartment.id, Decimal("-500"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x"
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_vacant_apartment_debt_is_in_the_snapshot(self):
        data = build_portal_data(self.user)
        resident = next(row for row in data["residents"] if row["apartmentBackendId"] == self.apartment.id)
        self.assertEqual(resident["name"], "Unassigned owner")
        self.assertEqual(resident["balance"], -500.0)
        self.assertEqual(resident["status"], "debtor")
        self.assertEqual(data["transactions"][0]["residentId"], resident["id"])
        self.assertIsNone(data["transactions"][0]["ownerBackendId"])

    def test_new_owner_sees_the_apartment_account_history(self):
        owner = Owner.objects.create(fio="Next Owner", phone="+998901234567", apartment=self.apartment)
        self.assertEqual(owner.balance, Decimal("-500"))

        data = build_portal_data(self.user)
        resident = next(row for row in data["residents"] if row["apartmentBackendId"] == self.apartment.id)
        self.assertEqual((resident["name"], resident["balance"]), ("Next Owner", -500.0))
        self.assertEqual(data["transactions"][0]["residentId"], resident["id"])

        rows = self.client.get(reverse("portal:api_list_transactions"), {"owner_id": owner.id}).json()["results"]
        self.assertEqual([row["signedAmount"] for row in rows], [-500.0])


class PortalResidentSearchTests(TestCase):
    """The command search asks /api/lists/residents/?search= and must get only matching residents."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64A", complex=complex_obj, author="test")
        first = Apartment.objects.create(
            number="1", area=Decimal("50"), building=building, personal_account="1391677280", author="test"
        )
        second = Apartment.objects.create(
            number="2", area=Decimal("50"), building=building, personal_account="1383000001", author="test"
        )
        cls.azamat = Owner.objects.create(fio="Azamat Yaxyoyev", phone="+998901139933", apartment=first)
        cls.muhammad = Owner.objects.create(fio="Muhammad Yaxyoyev", phone="+998773487001", apartment=second)

    def setUp(self):
        self.client.force_login(self.user)

    def search(self, query):
        response = self.client.get(reverse("portal:api_list_residents"), {"search": query})
        return [row["name"] for row in response.json()["results"]]

    def test_search_returns_only_matching_residents(self):
        self.assertEqual(self.search("Aza"), ["Azamat Yaxyoyev"])
        self.assertEqual(self.search("aza"), ["Azamat Yaxyoyev"])

    def test_search_matches_phone_apartment_and_shared_surname(self):
        self.assertEqual(self.search("487001"), ["Muhammad Yaxyoyev"])
        self.assertEqual(self.search("Yaxyoyev"), ["Azamat Yaxyoyev", "Muhammad Yaxyoyev"])
        self.assertEqual(self.search("64A"), ["Azamat Yaxyoyev", "Muhammad Yaxyoyev"])

    def test_search_with_no_match_returns_nothing(self):
        self.assertEqual(self.search("Zzz"), [])

    def test_search_and_sort_by_personal_account(self):
        self.assertEqual(self.search("1391677280"), ["Azamat Yaxyoyev"])
        self.assertEqual(self.search("677280"), ["Azamat Yaxyoyev"])

        response = self.client.get(reverse("portal:api_list_residents"), {"ordering": "personal_account"})

        rows = response.json()["results"]
        self.assertEqual([row["personalAccount"] for row in rows], ["1383000001", "1391677280"])


class PortalBoundedSnapshotTests(TestCase):
    """The snapshot keeps a few rows per account; whole-history figures come pre-aggregated."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64", complex=complex_obj, author="test")
        apartment = Apartment.objects.create(number="1", area=Decimal("50.00"), building=building, author="test")
        cls.owner = Owner.objects.create(fio="Long History", phone="+998901234567", apartment=apartment)
        for index in range(12):
            apply_balance_change(cls.owner.id, Decimal(1000 + index), Transaction.TYPE_PAYME)
        apply_balance_change(cls.owner.id, Decimal("-300"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")

    def setUp(self):
        self.client.force_login(self.user)

    def test_snapshot_keeps_only_the_latest_rows_of_each_account(self):
        data = build_portal_data(self.user)
        newest = list(Transaction.objects.order_by("-created_at", "-id").values_list("id", flat=True)[:5])
        self.assertEqual([row["backendId"] for row in data["transactions"]], newest)

    def test_payment_summary_covers_the_whole_history(self):
        summary = build_portal_data(self.user)["paymentSummary"]
        self.assertEqual(summary["channels"]["Payme"], {"amount": float(sum(1000 + i for i in range(12))), "residents": 1})
        self.assertEqual(sum(row["count"] for row in summary["monthly"]), 12)
        self.assertEqual({row["channel"] for row in summary["monthly"]}, {"Payme"})
        self.assertEqual(summary["currentMonthPayments"], 12)

    def test_balance_summary_gives_house_balances_at_the_end_of_each_month(self):
        # A charge two months back, then this month's payments: the old month shows the debt it had then.
        apartment = self.owner.apartment
        today = timezone.localtime()
        old = (today.replace(day=1) - timedelta(days=40)).replace(day=15)
        charge = apply_account_change(apartment.pk, Decimal("-50000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        Transaction.objects.filter(pk=charge.pk).update(created_at=old)
        apartment.refresh_from_db()

        data = build_portal_data(self.user)
        months = data["paymentSummary"]["months"]
        current, old_key = today.strftime("%Y-%m"), old.strftime("%Y-%m")
        self.assertIn(current, months)
        rows = {row["month"]: row for row in data["balanceSummary"]}
        self.assertEqual(set(rows), set(months))
        # Now: this month's payments (+11 766) against the old charge.
        self.assertEqual(apartment.balance, Decimal("-38234.00"))
        self.assertEqual((rows[current]["debt"], rows[current]["overpayment"], rows[current]["debtors"]), (38234.0, 0.0, 1))
        # At the end of the old month only the charge had been posted.
        self.assertEqual((rows[old_key]["debt"], rows[old_key]["overpayment"], rows[old_key]["debtors"]), (50000.0, 0.0, 1))

    def test_debtors_list_follows_the_statistics_month_and_house(self):
        # Another house: a resident charged this month and a vacant apartment in debt.
        other = Building.objects.create(number="61", complex=self.owner.apartment.building.complex, author="test")
        flat = Apartment.objects.create(number="7", area=Decimal("40.00"), building=other, author="test")
        Owner.objects.create(fio="New Debtor", phone="+998900000007", apartment=flat)
        apply_account_change(flat.pk, Decimal("-80000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        vacant = Apartment.objects.create(number="8", area=Decimal("40.00"), building=other, author="test")
        apply_account_change(vacant.pk, Decimal("-10000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        # The first house's resident: charged two months ago, paid in full since (setUpTestData payments).
        today = timezone.localtime()
        old = (today.replace(day=1) - timedelta(days=40)).replace(day=15)
        charge = apply_account_change(self.owner.apartment_id, Decimal("-10000"), Transaction.TYPE_CHARGE, description="[billing:period-charge] x")
        Transaction.objects.filter(pk=charge.pk).update(created_at=old)

        def debtors(**params):
            response = self.client.get(reverse("portal:api_list_debtors"), params)
            self.assertEqual(response.status_code, 200)
            return [(row["name"], row["house"], row["apartment"], row["debt"]) for row in response.json()["results"]]

        now = [("New Debtor", "61", "7", 80000.0), ("Unassigned owner", "61", "8", 10000.0)]
        self.assertEqual(debtors(month=today.strftime("%Y-%m")), now)
        self.assertEqual(debtors(), now)
        self.assertEqual(debtors(month=today.strftime("%Y-%m"), building_id=other.pk), now)
        self.assertEqual(debtors(month=today.strftime("%Y-%m"), building_id=self.owner.apartment.building_id), [])
        # At the end of the old month only the old charge existed.
        self.assertEqual(debtors(month=old.strftime("%Y-%m")), [("Long History", "64", "1", 10000.0)])

    def test_resident_collected_total_uses_the_whole_history(self):
        complex_row = build_portal_data(self.user)["complexes"][0]
        self.assertEqual(complex_row["baselineCollected"], float(sum(1000 + i for i in range(12))))

    def test_list_endpoint_returns_the_whole_account_history(self):
        response = self.client.get(
            reverse("portal:api_list_transactions"), {"apartment_id": self.owner.apartment_id, "page_size": 100}
        )
        self.assertEqual(response.json()["total"], 13)


class PortalBillingNoticeTests(TestCase):
    """Payment notice: one text for the window preview, the PDF and Telegram."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(
            username="desk", password="unused-password", first_name="Dilnoza", last_name="Karimova"
        )
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        building = Building.objects.create(number="64A", complex=complex_obj, author="test")
        cls.apartment = Apartment.objects.create(
            number="12", area=Decimal("50"), building=building, personal_account="1383364304", author="test"
        )
        cls.owner = Owner.objects.create(fio="Aziz Karimov", phone="+998901112233", apartment=cls.apartment)

    def setUp(self):
        self.client.force_login(self.user)

    def test_preview_has_the_facts_legal_basis_and_signatures(self):
        apply_account_change(self.apartment.pk, Decimal("-150000"), Transaction.TYPE_MANUAL)

        response = self.client.get(
            reverse("portal:api_billing_notice_preview"), {"owner_id": self.owner.pk, "language": "uz"}
        )

        self.assertEqual(response.status_code, 200)
        notice = response.json()["notice"]
        facts = {field["key"]: field["value"] for field in notice["fields"]}
        self.assertEqual(
            (facts["sector"], facts["house"], facts["apartment"], facts["resident"], facts["phone"], facts["debt"]),
            ("Mirabad Avenue", "64A", "12", "Aziz Karimov", "+998901112233", "150 000 so'm"),
        )
        self.assertEqual(facts["date"], timezone.localdate().strftime("%d.%m.%Y"))
        self.assertEqual(facts["account"], "1383364304")
        text = " ".join(notice["paragraphs"])
        for reference in ("132-moddasi", "134-moddasiga", "O'RQ-581", "236 va 237-moddalariga", "140 va 141-moddalariga"):
            self.assertIn(reference, text)
        signatures = {item["key"]: item["value"] for item in notice["signatures"]}
        self.assertEqual(signatures["operator"], "Dilnoza Karimova")
        self.assertTrue(signatures["director"])

    def test_pdf_download_and_no_debt_refusal(self):
        url = reverse("portal:api_billing_notice_pdf")
        refused = self.client.post(url, {"owner_id": self.owner.pk, "language": "ru"}, content_type="application/json")
        self.assertEqual(refused.status_code, 409)

        apply_account_change(self.apartment.pk, Decimal("-150000"), Transaction.TYPE_MANUAL)
        response = self.client.post(url, {"owner_id": self.owner.pk, "language": "ru"}, content_type="application/json")

        self.assertEqual(response.status_code, 200)
        download = response.json()["download"]
        self.assertEqual(download["mime"], "application/pdf")
        self.assertTrue(base64.b64decode(download["content"]).startswith(b"%PDF"))
        self.assertTrue(download["filename"].endswith(".pdf"))
        # Notices are not money: the activity log does not keep them.
        self.assertFalse(AuditEvent.objects.exists())

    def test_unknown_resident(self):
        response = self.client.get(reverse("portal:api_billing_notice_preview"), {"owner_id": 999999})
        self.assertEqual(response.status_code, 404)


class PortalBillingPeriodPagesTests(TestCase):
    """/billing/periods: the period list with its Create window, and a period's Excel-like house sheets."""

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        cls.house_60 = Building.objects.create(number="60", complex=complex_obj, author="test")
        cls.house_60v = Building.objects.create(number="60V", complex=complex_obj, author="test")
        cls.first = Apartment.objects.create(
            number="1", area=Decimal("171.94"), building=cls.house_60, personal_account="1391677280", author="test"
        )
        cls.rest = Apartment.objects.create(number="2", area=Decimal("173021.18"), building=cls.house_60, author="test")
        cls.other = Apartment.objects.create(number="1", area=Decimal("50.00"), building=cls.house_60v, author="test")
        Owner.objects.create(fio="Aliyev Jasur", phone="+998901234567", apartment=cls.first)
        Owner.objects.create(fio="Karimova Dilnoza", phone="+998901234568", apartment=cls.rest)

    def setUp(self):
        self.client.force_login(self.user)

    def _create(self, name="2025-Dekabr", start="2025-12-01", end="2025-12-31", follow=False):
        return self.client.post(
            reverse("portal:billing_periods"), {"name": name, "start_date": start, "end_date": end}, follow=follow
        )

    def _sheet(self, period, **data):
        data.setdefault("building", self.house_60.pk)
        data.setdefault("sheet", "heating")
        return self.client.post(reverse("portal:billing_period_sheet", args=[period.pk]), data, follow=True)

    def _messages(self, response):
        return [str(message) for message in response.context["messages"]]

    def _house_60_december(self):
        """The operators' example: house 60, December; apartment 1 heated 30 of 31 days."""
        self._create()
        period = BillingPeriod.objects.get(name="2025-Dekabr")
        first_heating = HeatingRecord.objects.get(period=period, apartment=self.first)
        rest_heating = HeatingRecord.objects.get(period=period, apartment=self.rest)
        other_heating = HeatingRecord.objects.get(period=period, apartment=self.other)
        self._sheet(
            period, action="save_heating", row=[first_heating.pk, rest_heating.pk],
            **{
                f"from-{first_heating.pk}": "2025-12-01", f"until-{first_heating.pk}": "2025-12-31", f"days-{first_heating.pk}": "30",
                f"from-{rest_heating.pk}": "2025-12-01", f"until-{rest_heating.pk}": "2025-12-31", f"days-{rest_heating.pk}": "1",
            },
        )
        self._sheet(
            period, building=self.house_60v.pk, action="save_heating", row=[other_heating.pk],
            **{f"from-{other_heating.pk}": "", f"until-{other_heating.pk}": "", f"days-{other_heating.pk}": ""},
        )
        self._sheet(
            period, action="save_expenses",
            new_name=["Обслуживание котельной", "Холодная вода"], new_amount=["5 967 950", "1 060 800"],
        )
        self._sheet(period, action="save_gas", new_gstart=["579 702,44"], new_gend=["601 719,18"], new_gprice=["1 000"])
        self._sheet(
            period, sheet="hot_water", action="save_expenses",
            new_name=["Обслуживание котельной", "Электроэнергия"], new_amount=["5967950", "1500000"],
        )
        self._sheet(period, sheet="hot_water", action="save_gas", new_gstart=["1000"], new_gend=["5893"], new_gprice=["1000"])
        readings = {
            reading.apartment_id: reading.pk
            for reading in HotWaterMeterReading.objects.filter(period=period, apartment__building=self.house_60)
        }
        self._sheet(
            period, sheet="hot_water", action="save_hot_water", row=list(readings.values()),
            **{
                f"start-{readings[self.first.pk]}": "89,1", f"end-{readings[self.first.pk]}": "100",
                f"start-{readings[self.rest.pk]}": "0", f"end-{readings[self.rest.pk]}": "214,3",
            },
        )
        return period

    def test_create_makes_an_open_period_with_both_tables_for_every_house(self):
        response = self._create(follow=True)

        period = BillingPeriod.objects.get(name="2025-Dekabr")
        self.assertRedirects(response, reverse("portal:billing_periods"))
        self.assertEqual(period.status, BillingPeriod.STATUS_DRAFT)
        self.assertEqual(set(period.buildings.all()), {self.house_60, self.house_60v})
        heating = HeatingRecord.objects.filter(period=period)
        self.assertEqual({row.apartment_id for row in heating}, {self.first.pk, self.rest.pk, self.other.pk})
        self.assertEqual({(row.heated_from, row.heated_until) for row in heating}, {(date(2025, 12, 1), None)})
        self.assertEqual(HotWaterMeterReading.objects.filter(period=period).count(), 3)
        self.assertTrue(AuditEvent.objects.filter(title="Billing period created").exists())
        self.assertContains(response, "2025-Dekabr")

    def test_create_refuses_overlapping_or_reversed_dates(self):
        self._create()
        overlapping = self._create(name="Again", start="2025-12-20", end="2026-01-19")
        reversed_dates = self._create(name="Backwards", start="2026-02-10", end="2026-02-01")

        self.assertEqual(BillingPeriod.objects.count(), 1)
        self.assertContains(overlapping, "boshqa davr bilan kesishadi")
        self.assertContains(reversed_dates, "boshlanish sanasidan oldin")
        self.assertTrue(reversed_dates.context["open_create"])  # the window stays open with what was typed

    def test_list_shows_periods_in_order_and_opens_on_the_latest(self):
        for month in range(1, 13):
            BillingPeriod.objects.create(name=f"2026-{month:02d}", start_date=date(2026, month, 1), end_date=date(2026, month, 1))

        latest = self.client.get(reverse("portal:billing_periods"))
        first_page = self.client.get(reverse("portal:billing_periods"), {"page": 1})

        self.assertEqual([row["period"].name for row in latest.context["rows"]], ["2026-11", "2026-12"])
        self.assertEqual([row["period"].name for row in first_page.context["rows"]][:3], ["2026-01", "2026-02", "2026-03"])
        self.assertContains(latest, reverse("portal:billing_period_sheet", args=[BillingPeriod.objects.get(name="2026-12").pk]))

    def test_old_address_opens_the_periods_page(self):
        response = self.client.get("/billing/periods/create")
        self.assertRedirects(response, reverse("portal:billing_periods"))

    def test_house_60_december_matches_the_excel_sheets(self):
        period = self._house_60_december()

        heating = self.client.get(reverse("portal:billing_period_sheet", args=[period.pk]), {"building": self.house_60.pk})
        rows = {row["apt"]: row for row in heating.context["rows"]}
        self.assertEqual((rows["1"]["range_days"], rows["1"]["days"], rows["1"]["heated_area"]), (31, "30", "5 158,2"))
        self.assertEqual(rows["1"]["amount"], "840 851,77")
        self.assertEqual(heating.context["total_cost"], "29 045 490")  # gas 22 016 740 + boiler 5 967 950 + cold water 1 060 800
        self.assertEqual(heating.context["total_volume"], "178 179,38")
        self.assertEqual(heating.context["tariff"], "163,0126")

        hot_water = self.client.get(
            reverse("portal:billing_period_sheet", args=[period.pk]), {"building": self.house_60.pk, "sheet": "hot_water"}
        )
        rows = {row["apt"]: row for row in hot_water.context["rows"]}
        self.assertEqual((rows["1"]["consumption"], rows["1"]["amount"]), ("10,9", "598 287,54"))
        self.assertEqual((hot_water.context["total_volume"], hot_water.context["total_cost"]), ("225,2", "12 360 950"))

        response = self._sheet(period, action="set_status", status="closed")
        period.refresh_from_db()
        self.first.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_CLOSED)
        self.assertEqual(self.first.balance, Decimal("-1439139.31"))
        self.assertTrue(any("yopildi" in message for message in self._messages(response)))

    def test_heating_errors_keep_what_was_typed(self):
        self._create()
        period = BillingPeriod.objects.get()
        record = HeatingRecord.objects.get(period=period, apartment=self.first)

        response = self._sheet(
            period, action="save_heating", row=[record.pk],
            **{f"from-{record.pk}": "2025-12-11", f"until-{record.pk}": "", f"days-{record.pk}": "25"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(any("kv 1" in message and "ko'p" in message for message in self._messages(response)))
        self.assertEqual({row["apt"]: row["days"] for row in response.context["rows"]}["1"], "25")
        record.refresh_from_db()
        self.assertEqual((record.heated_from, record.days_given), (date(2025, 12, 1), None))

    def test_closed_period_is_read_only_until_reopened_and_closed_again(self):
        period = self._house_60_december()
        self._sheet(period, action="set_status", status="closed")
        record = HeatingRecord.objects.get(period=period, apartment=self.first)
        days_31 = {f"from-{record.pk}": "2025-12-01", f"until-{record.pk}": "", f"days-{record.pk}": "31"}

        refused = self._sheet(period, action="save_heating", row=[record.pk], **days_31)
        self.assertTrue(any("yopiq" in message for message in self._messages(refused)))
        record.refresh_from_db()
        self.assertEqual(record.days_given, 30)

        self._sheet(period, action="set_status", status="open")
        period.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_REOPENED)
        self._sheet(period, action="save_heating", row=[record.pk], **days_31)
        self._sheet(period, action="set_status", status="closed")

        period.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_CLOSED)
        self.assertTrue(AuditEvent.objects.filter(title="Billing period recalculated").exists())
        heating = Charge.objects.get(period=period, apartment=self.first, charge_type="heating", is_cancelled=False)
        self.assertGreater(heating.amount, Decimal("840851.77"))  # 31 days now instead of 30

    def test_save_with_status_closed_closes_the_period_and_charges_the_accounts(self):
        period = self._house_60_december()
        record = HeatingRecord.objects.get(period=period, apartment=self.first)
        self.rest.refresh_from_db()
        rest_before = self.rest.balance

        response = self._sheet(
            period, action="save_heating", row=[record.pk], status="closed",
            **{f"from-{record.pk}": "2025-12-01", f"until-{record.pk}": "", f"days-{record.pk}": "30"},
        )

        period.refresh_from_db()
        self.first.refresh_from_db()
        self.rest.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_CLOSED)
        self.assertEqual(self.first.balance, Decimal("-1439139.31"))  # nothing paid before: all of it is debt
        self.assertLess(self.rest.balance, rest_before)
        self.assertTrue(any("yopildi" in message for message in self._messages(response)))

    def test_save_with_status_closed_keeps_the_table_when_the_period_cannot_close(self):
        self._create()
        period = BillingPeriod.objects.get()
        record = HeatingRecord.objects.get(period=period, apartment=self.first)

        response = self._sheet(
            period, action="save_heating", row=[record.pk], status="closed",
            **{f"from-{record.pk}": "2025-12-01", f"until-{record.pk}": "", f"days-{record.pk}": "28"},
        )

        period.refresh_from_db()
        record.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_DRAFT)  # no costs entered yet
        self.assertEqual(record.days_given, 28)
        self.assertTrue(any("Davr yopilmadi" in message for message in self._messages(response)))

    def test_paid_ahead_balance_covers_the_charge(self):
        period = self._house_60_december()
        apply_account_change(self.first.pk, Decimal("2000000"), Transaction.TYPE_CASH, description="Paid ahead")

        self._sheet(period, action="set_status", status="closed")

        self.first.refresh_from_db()
        self.assertEqual(self.first.balance, Decimal("560860.69"))  # 2 000 000 − 1 439 139.31: still in credit

    def test_delete_removes_an_open_period_but_never_a_billed_one(self):
        billed = self._house_60_december()
        self._sheet(billed, action="set_status", status="closed")
        self._create(name="2026-Yanvar", start="2026-01-01", end="2026-01-31")
        draft = BillingPeriod.objects.get(name="2026-Yanvar")

        refused = self._sheet(billed, action="delete")
        deleted = self._sheet(draft, action="delete")

        self.assertTrue(BillingPeriod.objects.filter(pk=billed.pk).exists())
        self.assertTrue(any("o'chirib bo'lmaydi" in message for message in self._messages(refused)))
        self.assertFalse(BillingPeriod.objects.filter(pk=draft.pk).exists())
        self.assertFalse(HeatingRecord.objects.filter(period_id=draft.pk).exists())
        self.assertRedirects(deleted, reverse("portal:billing_periods"))

    def test_refresh_adds_apartments_and_houses_created_after_the_period(self):
        self._create()
        period = BillingPeriod.objects.get()
        new_apartment = Apartment.objects.create(number="3", area=Decimal("60.00"), building=self.house_60, author="test")
        new_house = Building.objects.create(number="62", complex=self.house_60.complex, author="test")
        Apartment.objects.create(number="1", area=Decimal("70.00"), building=new_house, author="test")

        self._sheet(period, action="refresh")

        self.assertIn(new_house.pk, period.buildings.values_list("pk", flat=True))
        self.assertTrue(HeatingRecord.objects.filter(period=period, apartment=new_apartment, heated_from=date(2025, 12, 1)).exists())
        self.assertEqual(HotWaterMeterReading.objects.filter(period=period).count(), 5)

    def test_export_has_a_heating_and_a_hot_water_sheet_per_house(self):
        period = self._house_60_december()

        response = self.client.get(reverse("portal:billing_period_export", args=[period.pk]), {"lang": "ru"})

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["60 Отопление", "60 ГВС", "60V Отопление", "60V ГВС"])
        numbers = [
            Decimal(str(cell))
            for row in workbook["60 Отопление"].iter_rows(values_only=True)
            for cell in row
            if isinstance(cell, (int, float, Decimal)) and not isinstance(cell, bool)
        ]
        self.assertIn(Decimal("840851.77"), numbers)
        for title in ("60 Отопление", "60 ГВС"):
            rows = list(workbook[title].iter_rows(values_only=True))
            self.assertEqual(rows[3][:4], ("№", "Ф.И.О.", "Лицевой счет", "Кв"))
            self.assertEqual(rows[4][2], "1391677280")

    def test_sheet_shows_the_personal_account(self):
        period = self._house_60_december()

        for sheet in ("heating", "hot_water"):
            response = self.client.get(
                reverse("portal:billing_period_sheet", args=[period.pk]), {"building": self.house_60.pk, "sheet": sheet}
            )
            self.assertContains(response, 'data-pt="personal_account"')
            self.assertContains(response, '<td class="bp-account">1391677280</td>', html=False)


class PortalInputValidationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="unused-password")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        cls.building = Building.objects.create(number="64", complex=complex_obj, author="test")
        cls.apartment = Apartment.objects.create(number="1", area=Decimal("50.00"), building=cls.building, author="test")
        cls.owner = Owner.objects.create(fio="Owner", phone="+998901234567", apartment=cls.apartment)

    def setUp(self):
        self.client.force_login(self.user)

    def _post(self, name, payload):
        return self.client.post(reverse(name), json.dumps(payload), content_type="application/json")

    def test_payment_rejects_non_finite_and_oversized_amounts(self):
        for amount in ["NaN", "Infinity", "-Infinity", "1e15", "0.001"]:
            response = self._post(
                "portal:api_create_payment",
                {"owner_id": self.owner.id, "amount": amount, "operation": "credit", "channel": "cash"},
            )
            self.assertEqual(response.status_code, 400, amount)
        self.assertFalse(Transaction.objects.exists())

    def test_payment_amount_is_rounded_to_tiyin(self):
        response = self._post(
            "portal:api_create_payment",
            {"owner_id": self.owner.id, "amount": "100.005", "operation": "credit", "channel": "cash"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Transaction.objects.get().amount, Decimal("100.01"))

    def test_apartment_rejects_non_finite_area(self):
        for area in ["nan", "inf", "abc"]:
            response = self._post(
                "portal:api_create_apartment", {"building_id": self.building.id, "number": f"x-{area}", "area": area}
            )
            self.assertEqual(response.status_code, 400, area)

    def test_bad_counters_do_not_crash(self):
        response = self._post("portal:api_resident_kit_action", {"action": "contact_sheet", "visible_count": "many"})
        self.assertEqual(response.status_code, 200)
        response = self._post("portal:api_assign_status", {"context": "transactions", "status": "x", "selected": ["bad"]})
        self.assertEqual(response.json()["updated"], 0)

    def test_hot_water_total_is_consumption_not_raw_meter_readings(self):
        period = BillingPeriod.objects.create(name="April", start_date=date(2026, 4, 1), end_date=date(2026, 4, 30))
        HotWaterMeterReading.objects.create(
            period=period, apartment_id=self.apartment.pk, start_reading=Decimal("1200"), end_reading=Decimal("1205.5")
        )
        complex_row = build_portal_data(self.user)["complexes"][0]
        self.assertEqual(complex_row["waterM3"], 5.5)


class OperatorRightsTests(TestCase):
    """Only operators use the platform, and each of them has full rights in the portal and admin."""

    def test_a_new_account_is_staff_and_superuser(self):
        user = get_user_model().objects.create_user(username="kassir3", password="unused-password")

        self.assertEqual((user.is_staff, user.is_superuser), (True, True))
        # Rights narrowed later by hand stay narrowed.
        user.is_superuser = False
        user.save()
        user.refresh_from_db()
        self.assertFalse(user.is_superuser)

    def test_migration_gives_every_active_account_full_rights(self):
        from importlib import import_module

        from django.apps import apps as django_apps

        give = import_module("portal.migrations.0006_operators_full_access").give_operators_full_rights
        User = get_user_model()
        active = User.objects.create_user(username="old-operator", password="unused-password")
        gone = User.objects.create_user(username="left", password="unused-password", is_active=False)
        User.objects.update(is_staff=False, is_superuser=False)

        give(django_apps, None)

        active.refresh_from_db()
        gone.refresh_from_db()
        self.assertEqual((active.is_staff, active.is_superuser), (True, True))
        self.assertEqual((gone.is_staff, gone.is_superuser), (False, False))


class OperatorActivityLogTests(TestCase):
    """The log keeps payments, balance changes, billing period work and Telegram links, each with its operator."""

    password = "operator-test-password"

    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.operator = User.objects.create_user(username="kassir1", password=cls.password, first_name="Ali", last_name="Valiyev")
        cls.other_operator = User.objects.create_user(username="kassir2", password=cls.password)
        cls.head_admin = User.objects.create_superuser(username="boss", password=cls.password)
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        cls.building = Building.objects.create(number="64", complex=complex_obj, author="test")
        cls.apartment = Apartment.objects.create(number="5", area=Decimal("50.00"), building=cls.building, author="test")
        cls.owner = Owner.objects.create(fio="Mijoz", phone="+998901234567", apartment=cls.apartment)

    def _post_json(self, name, payload, **extra):
        return self.client.post(reverse(name), json.dumps(payload), content_type="application/json", **extra)

    def _pay(self, amount="250000", operation="credit", **extra):
        return self._post_json(
            "portal:api_create_payment",
            {"owner_id": self.owner.id, "amount": amount, "operation": operation, "channel": "cash"},
            **extra,
        )

    def test_cash_payment_is_logged_with_the_operator_account_and_ip(self):
        self.client.force_login(self.operator)
        response = self._pay(REMOTE_ADDR="10.0.0.7")
        self.assertEqual(response.status_code, 200)
        event = AuditEvent.objects.get()
        self.assertEqual(event.event_type, AuditEvent.TYPE_OPERATOR_PAYMENT)
        self.assertEqual((event.user, event.actor, event.ip_address), (self.operator, "kassir1", "10.0.0.7"))
        self.assertEqual(event.owner_id, self.owner.id)
        self.assertEqual(event.metadata["transaction_id"], Transaction.objects.get().id)
        self.assertEqual(Transaction.objects.get().created_by, "kassir1")

    def test_a_subtraction_is_a_balance_change(self):
        self.client.force_login(self.operator)
        self._pay(amount="50000", operation="debit")
        self.assertEqual(AuditEvent.objects.get().event_type, AuditEvent.TYPE_BALANCE)

    def test_sign_ins_admin_edits_notices_and_exports_are_not_logged(self):
        apply_account_change(self.apartment.pk, Decimal("-150000"), Transaction.TYPE_MANUAL)
        self.client.post(reverse("portal:login"), {"username": "kassir1", "password": "wrong"})
        self.client.post(reverse("portal:login"), {"username": "kassir1", "password": self.password})
        self._post_json("portal:api_billing_notice_pdf", {"owner_id": self.owner.id, "language": "uz"})
        self._post_json("portal:api_export_audit", {})
        self.client.post(reverse("portal:api_logout"))
        group = Group.objects.create(name="Kassa")
        self.client.force_login(self.head_admin)
        self.client.post("/admin/auth/group/", {"action": "delete_selected", "_selected_action": [group.pk], "post": "yes"})
        self.client.post(reverse("admin:auth_group_add"), {"name": "Arxiv"})

        self.assertFalse(Group.objects.filter(name="Kassa").exists())
        self.assertTrue(Group.objects.filter(name="Arxiv").exists())
        self.assertFalse(AuditEvent.objects.exists())

    def test_other_kinds_are_refused(self):
        with self.assertRaises(ValueError):
            record_event(actor="test", event_type="auth", title="Signed in", message="-")

    def test_billing_period_work_is_logged_as_edits_and_status_changes(self):
        from billing.models import BillingPeriod

        self.client.force_login(self.operator)
        self.client.post(reverse("portal:billing_periods"), {"name": "Aprel", "start_date": "2026-04-01", "end_date": "2026-04-30"})
        period = BillingPeriod.objects.get(name="Aprel")
        url = reverse("portal:billing_period_sheet", args=[period.pk])
        base = {"building": self.building.id, "sheet": "heating"}
        self.client.post(url, {**base, "action": "save_expenses", "new_name": "Qozonxona", "new_amount": "1500"})
        record = period.heating_records.get(apartment=self.apartment)
        self.client.post(url, {
            **base, "action": "save_heating", "row": record.pk,
            f"from-{record.pk}": "2026-04-01", f"until-{record.pk}": "", f"days-{record.pk}": "29",
        })
        self.client.post(url, {**base, "action": "set_status", "status": "closed"})
        self.client.post(url, {**base, "action": "set_status", "status": "open"})

        events = list(AuditEvent.objects.order_by("id").values_list("event_type", "title"))
        self.assertEqual(
            events,
            [
                (AuditEvent.TYPE_PERIOD_EDIT, "Billing period created"),
                (AuditEvent.TYPE_PERIOD_EDIT, "Expenses saved"),
                (AuditEvent.TYPE_PERIOD_EDIT, "Heating table saved"),
                (AuditEvent.TYPE_PERIOD_STATUS, "Billing period closed"),
                (AuditEvent.TYPE_PERIOD_STATUS, "Billing period reopened"),
            ],
        )
        self.assertTrue(all(event.user == self.operator for event in AuditEvent.objects.all()))
        self.assertEqual(AuditEvent.objects.get(title="Heating table saved").building_id, self.building.id)

    def test_failed_billing_action_leaves_no_log_entry(self):
        from billing.models import BillingPeriod

        self.client.force_login(self.operator)
        period = BillingPeriod.objects.create(name="Bo'sh", start_date=date(2026, 4, 1), end_date=date(2026, 4, 30))
        # No buildings or costs: closing fails, so nothing happened and nothing is logged.
        self.client.post(reverse("portal:billing_period_sheet", args=[period.pk]), {"action": "set_status", "status": "closed"})
        period.refresh_from_db()
        self.assertEqual(period.status, BillingPeriod.STATUS_DRAFT)
        self.assertFalse(AuditEvent.objects.exists())

    def test_operators_read_the_log_unless_their_rights_were_narrowed(self):
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(reverse("portal:api_list_audit")).status_code, 200)
        self.assertIn('id="audit-drawer"', self.client.get(reverse("portal:dashboard")).content.decode())

        type(self.operator).objects.filter(pk=self.operator.pk).update(is_staff=False, is_superuser=False)
        self.operator.refresh_from_db()
        self.assertEqual(self.client.get(reverse("portal:api_list_audit")).status_code, 403)
        self.assertEqual(self._post_json("portal:api_export_audit", {}).status_code, 403)
        self.assertEqual(build_portal_data(self.operator)["auditEvents"], [])
        page = self.client.get(reverse("portal:dashboard")).content.decode()
        self.assertNotIn('id="audit-drawer"', page)
        self.assertEqual(self.client.get("/admin/portal/auditevent/").status_code, 302)

    def test_the_log_is_filtered_by_operator_and_category(self):
        self.client.force_login(self.operator)
        self._pay(amount="100")
        self._pay(amount="40", operation="debit")
        self.client.force_login(self.other_operator)
        self._pay(amount="200")

        self.client.force_login(self.head_admin)
        page = self.client.get(reverse("portal:dashboard")).content.decode()
        self.assertIn('data-audit-filter="operator_payment"', page)
        self.assertIn("Ali Valiyev (kassir1)", page)
        rows = self.client.get(
            reverse("portal:api_list_audit"), {"user_id": self.operator.id, "type": AuditEvent.TYPE_OPERATOR_PAYMENT}
        ).json()["results"]
        self.assertEqual([(row["actor"], row["label"]) for row in rows], [("kassir1", "Payment: operator")])
        self.assertIn("100", rows[0]["message"])
        balance = self.client.get(reverse("portal:api_list_audit"), {"type": AuditEvent.TYPE_BALANCE}).json()["results"]
        self.assertEqual([row["actor"] for row in balance], ["kassir1"])

    def test_log_is_read_only_in_the_admin_and_split_into_categories(self):
        self.client.force_login(self.operator)
        self._pay()
        event = AuditEvent.objects.get()
        listing = self.client.get("/admin/portal/auditevent/", {"event_type__exact": AuditEvent.TYPE_OPERATOR_PAYMENT})
        self.assertContains(listing, "Resident balance updated")
        self.assertContains(listing, "Payment: Click / Payme")  # the category filter
        other = self.client.get("/admin/portal/auditevent/", {"event_type__exact": AuditEvent.TYPE_TELEGRAM})
        self.assertNotContains(other, "Resident balance updated")
        edit = self.client.post(f"/admin/portal/auditevent/{event.id}/change/", {"title": "edited"})
        delete = self.client.post(f"/admin/portal/auditevent/{event.id}/delete/", {"post": "yes"})
        self.assertEqual((edit.status_code, delete.status_code), (403, 403))
        event.refresh_from_db()
        self.assertEqual(event.title, "Resident balance updated")

    def test_deleting_an_operator_keeps_their_history(self):
        self.client.force_login(self.other_operator)
        self._pay(amount="777")
        self.other_operator.delete()
        event = AuditEvent.objects.get()
        self.assertEqual((event.user_id, event.actor), (None, "kassir2"))

    def test_migration_sorts_old_rows_into_categories_and_drops_the_rest(self):
        from importlib import import_module

        from django.apps import apps as django_apps

        categorize = import_module("portal.migrations.0007_audit_categories").categorize
        old = {
            ("balance", "Resident balance updated", "credit"): AuditEvent.TYPE_OPERATOR_PAYMENT,
            ("balance", "Resident balance updated", "debit"): AuditEvent.TYPE_BALANCE,
            ("balance", "Click payment received", "click"): AuditEvent.TYPE_ONLINE_PAYMENT,
            ("billing", "Heating table saved", ""): AuditEvent.TYPE_PERIOD_EDIT,
            ("billing", "Billing period closed", ""): AuditEvent.TYPE_PERIOD_STATUS,
            ("status", "Telegram connected", ""): AuditEvent.TYPE_TELEGRAM,
            ("status", "Telegram registration: phone not found", ""): None,
            ("auth", "Signed in", ""): None,
            ("admin", "Admin: House added", ""): None,
            ("export", "Portal export generated", ""): None,
        }
        ids = {}
        for (kind, title, detail), expected in old.items():
            metadata = {"source": "click"} if detail == "click" else {"operation": detail} if detail else {}
            event = AuditEvent.objects.create(event_type=kind, title=title, message="-", metadata=metadata)
            ids[event.pk] = expected

        categorize(django_apps, None)

        kept = dict(AuditEvent.objects.values_list("pk", "event_type"))
        self.assertEqual(kept, {pk: expected for pk, expected in ids.items() if expected})


class PortalPagesAfterOperatorsReviewTests(TestCase):
    """The operators' review of 2026-10-19: pages trimmed, operators per house, a Reports page, a small admin."""

    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.head_admin = User.objects.create_superuser(username="boss", password="unused-password")
        cls.operator = User.objects.create_user(username="kassir1", password="unused-password", first_name="Azamat")
        complex_obj = Complex.objects.create(title="Mirabad Avenue", address="", author="test")
        cls.building = Building.objects.create(number="64", complex=complex_obj, author="test", operator=cls.operator)
        Building.objects.create(number="64A", complex=complex_obj, author="test")

    def setUp(self):
        self.client.force_login(self.head_admin)

    def test_every_house_carries_its_operator(self):
        houses = {house["number"]: house for house in build_portal_data(self.head_admin)["complexes"][0]["buildingItems"]}
        self.assertEqual((houses["64"]["operatorId"], houses["64"]["operatorName"]), (self.operator.pk, "Azamat"))
        self.assertEqual((houses["64A"]["operatorId"], houses["64A"]["operatorName"]), (None, ""))
        page = self.client.get(reverse("portal:analytics")).content.decode()
        self.assertIn("data-analytics-operator-list", page)

    def test_pages_lost_the_tools_the_operators_crossed_out(self):
        payments = self.client.get(reverse("portal:billing")).content.decode()
        self.assertNotIn("Sector Revenue vs Debt", payments)
        self.assertIn("data-total-outstanding", payments)
        houses = self.client.get(reverse("portal:residential")).content.decode()
        self.assertNotIn("Export CSV", houses)
        self.assertIn('data-export-tool="off" data-reminders-tool="off" data-select-tool="off"', houses)
        residents = self.client.get(reverse("portal:dashboard")).content.decode()
        self.assertIn('data-stat-card="total-debtors" data-debtors-list-open', residents)
        self.assertNotIn("Generate Usage Report", residents)

    def test_reports_page_is_in_the_sidebar(self):
        page = self.client.get(reverse("portal:reports"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, reverse("portal:reports"))
        self.assertContains(page, 'data-export-context="usage-report"')

    def test_admin_keeps_only_what_the_portal_cannot_manage(self):
        from django.contrib import admin

        registered = {f"{model._meta.app_label}.{model.__name__}" for model in admin.site._registry}
        self.assertEqual(
            registered,
            {
                "auth.Group", "auth.User", "main_app.Apartment", "main_app.Building",
                "main_app.Complex", "main_app.Owner", "payments.ClickTransaction", "payments.PaymeTransaction",
                "payments.Transaction", "portal.AuditEvent", "portal.WorkspaceProfile",
            },
        )
        page = self.client.get(reverse("admin:main_app_building_change", args=[self.building.pk]))
        self.assertContains(page, 'name="operator"')

